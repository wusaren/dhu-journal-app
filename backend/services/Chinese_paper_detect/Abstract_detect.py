#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import sys
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

# 添加项目根目录到 sys.path 以支持独立运行
if __name__ == "__main__" and __package__ is None:
    file = Path(__file__).resolve()
    parent, root = file.parent, file.parents[1]
    sys.path.append(str(root))
    try:
        sys.path.remove(str(parent))
    except ValueError:
        pass

def should_skip_check(check_name):
    """
    判断是否应该跳过某个检测项
    参数:
        check_name: 检测项名称 (font_size, bold, italic, alignment, spacing, indent)
    返回:
        bool: True表示跳过该检测项，False表示执行该检测项
    """
    global _skip_checks_config
    if _skip_checks_config is None:
        return False
    return check_name in _skip_checks_config

# 全局变量，用于存储当前模块的跳过检测项配置
_skip_checks_config = []

"""
=== 论文格式检测系统 - 中文摘要检测器 ===

【中文摘要检测 (Chinese Abstract Detection)】

1. 【结构检测】
   - "摘 要"标题格式（中间一个空格）
   - 标题与内容之间隔一行
   - 摘要正文可分段
   - 内容长度验证（300-2000字符）

2. 【格式检测】
   - 标题格式：黑体3号（16pt），加粗，1倍行距，居中对齐，无缩进
   - 正文格式：宋体小四（12pt），英文Times New Roman，1.25倍行距，左对齐，首行缩进2字符

3. 【技术特性】
   - 支持中英文字体分别检测
   - 标题和正文分开检测
   - 智能段落定位（处理标题与正文之间的空行）
   - 完整的结构、段落、格式三维度检测
   - 详细的格式错误诊断和建议
"""

# ---------- 模板加载 ----------
def resolve_template_path(identifier):
    """解析模板路径，支持文件路径和模板名称"""
    if os.path.isfile(identifier):
        return identifier
    
    # 获取当前文件所在目录
    current_dir = Path(__file__).parent
    # 尝试多个可能的模板目录
    candidates = [
        os.path.join("templates", identifier + ".json"),
        os.path.join("Chinese_paper_detect_template", identifier + ".json"),
        os.path.join("services", "Chinese_paper_detect_template", identifier + ".json"),
        str(current_dir.parent / "Chinese_paper_detect_template" / (identifier + ".json")),
        str(current_dir / ".." / "Chinese_paper_detect_template" / (identifier + ".json")),
    ]
    
    # 添加绝对路径候选
    for candidate in candidates:
        abs_path = os.path.abspath(candidate)
        if os.path.isfile(abs_path):
            return abs_path
        if os.path.isfile(candidate):
            return candidate
    
    raise FileNotFoundError(f"Template not found: '{identifier}' (tried file path and {candidates})")

def load_template(identifier):
    """加载JSON模板文件"""
    tpl_path = resolve_template_path(identifier)
    with open(tpl_path, 'r', encoding='utf-8') as f:
        tpl = json.load(f)
    return tpl

# ---------- 字体检测函数（支持中英文字体分别检测）----------
def detect_font_for_run(run, paragraph=None):
    """
    检测run的字体信息，包括中英文字体
    返回: (font_size, font_ascii, font_eastasia, is_bold, is_italic, line_spacing)
    """
    font_size = None
    font_name_ascii = None
    font_name_eastasia = None
    is_bold = None
    is_italic = False
    
    if not run:
        return 12.0, "Times New Roman", "宋体", False, False, 1.0
    
    # 1. 检测字号
    try:
        # 直接格式
        if run.font and run.font.size and hasattr(run.font.size, 'pt'):
            font_size = float(run.font.size.pt)

        # 样式：run.style
        if font_size is None and getattr(run, "style", None) and getattr(run.style, "font", None):
            try:
                if run.style.font.size and hasattr(run.style.font.size, "pt"):
                    font_size = float(run.style.font.size.pt)
            except Exception:
                pass

        # 样式：paragraph.style
        if font_size is None and paragraph and paragraph.style and getattr(paragraph.style, "font", None):
            try:
                if paragraph.style.font.size and hasattr(paragraph.style.font.size, "pt"):
                    font_size = float(paragraph.style.font.size.pt)
            except Exception:
                pass

        # 样式 XML（段落样式的 rPr 或 run 的 rPr）
        if font_size is None and paragraph and paragraph.style and hasattr(paragraph.style, "element"):
            sz_nodes = paragraph.style.element.xpath('.//w:sz')
            if sz_nodes and sz_nodes[0].get(qn('w:val')):
                font_size = float(sz_nodes[0].get(qn('w:val'))) / 2.0

        # 直接从 run 的 rPr 读取
        if font_size is None and hasattr(run._element, 'rPr'):
            sz_nodes = run._element.xpath('.//w:sz')
            if sz_nodes and sz_nodes[0].get(qn('w:val')):
                font_size = float(sz_nodes[0].get(qn('w:val'))) / 2.0
    except Exception:
        pass
    
    font_size = font_size if font_size is not None else 12.0
    
    # 2. 检测英文字体（ASCII）和中文字体（EastAsia）
    try:
        # 优先从run.font.name读取
        if run.font and run.font.name:
            font_name_ascii = run.font.name
        
        # 从XML中读取w:rFonts
        if hasattr(run._element, 'rPr'):
            rpr = run._element.rPr
            if rpr is not None:
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is not None:
                    # 读取各种字体属性
                    xml_ascii = rfonts.get(qn('w:ascii'))
                    xml_hansi = rfonts.get(qn('w:hAnsi'))
                    xml_eastasia = rfonts.get(qn('w:eastAsia'))
                    
                    # ASCII字体（英文）
                    if xml_ascii:
                        font_name_ascii = xml_ascii
                    elif xml_hansi and font_name_ascii is None:
                        font_name_ascii = xml_hansi
                    
                    # EastAsia字体（中文）
                    if xml_eastasia:
                        font_name_eastasia = xml_eastasia
                    elif xml_hansi and font_name_eastasia is None:
                        font_name_eastasia = xml_hansi
    except Exception as e:
        print(f"  字体检测异常: {e}")
    
    font_name_ascii = font_name_ascii if font_name_ascii else "Times New Roman"
    font_name_eastasia = font_name_eastasia if font_name_eastasia else "宋体"
    
    # 3. 检测加粗和斜体
    try:
        if run.font:
            if run.font.bold is not None:
                is_bold = run.font.bold
            if run.font.italic is not None:
                is_italic = run.font.italic

        # 样式：run.style
        if is_bold is None and getattr(run, "style", None) and getattr(run.style, "font", None):
            try:
                if run.style.font.bold is not None:
                    is_bold = run.style.font.bold
            except Exception:
                pass

        # 样式：paragraph.style
        if is_bold is None and paragraph and paragraph.style and getattr(paragraph.style, "font", None):
            try:
                if paragraph.style.font.bold is not None:
                    is_bold = paragraph.style.font.bold
            except Exception:
                pass

        # 样式 XML：run 的 rPr 内的 <w:b>
        if is_bold is None and hasattr(run._element, 'rPr'):
            try:
                b_nodes = run._element.xpath('.//w:b')
                if b_nodes:
                    is_bold = True
            except Exception:
                pass

        is_bold = bool(is_bold) if is_bold is not None else False
        is_italic = bool(is_italic)
    except Exception:
        pass
    
    # 4. 行间距检测
    line_spacing = 1.0  # 默认单倍行距
    try:
        # 优先级1: 检查段落直接格式的行距
        if paragraph and paragraph.paragraph_format.line_spacing:
            line_spacing = float(paragraph.paragraph_format.line_spacing)
        # 优先级2: 检查段落样式的行距
        elif paragraph and paragraph.style and paragraph.style.paragraph_format.line_spacing:
            line_spacing = float(paragraph.style.paragraph_format.line_spacing)
    except Exception:
        pass
    
    return font_size, font_name_ascii, font_name_eastasia, is_bold, is_italic, line_spacing

def get_font_size(pt_size, tpl=None):
    """字体大小转换为中文字号"""
    if tpl and 'check_rules' in tpl and 'font_size_mapping' in tpl['check_rules']:
        size_config = tpl['check_rules']['font_size_mapping']
        size_map = {}
        for key, value in size_config.items():
            try:
                size_map[float(key)] = value
            except (ValueError, TypeError):
                continue
    else:
        size_map = {
            9: "小五", 10.5: "五号", 12: "小四", 14: "四号",
            16: "三号", 18: "小二", 22: "二号", 24: "小一", 26: "一号"
        }
    
    if not size_map:
        return f"{pt_size}pt"
    
    closest_size = min(size_map.keys(), key=lambda x: abs(x - pt_size))
    return size_map[closest_size]

def get_line_spacing_name(spacing, tpl=None):
    """行间距转换为中文描述"""
    if tpl and 'check_rules' in tpl and 'line_spacing_mapping' in tpl['check_rules']:
        spacing_config = tpl['check_rules']['line_spacing_mapping']
        spacing_map = {}
        for key, value in spacing_config.items():
            try:
                spacing_map[float(key)] = value
            except (ValueError, TypeError):
                continue
    else:
        spacing_map = {
            1.0: "单倍行距", 1.15: "1.15倍行距", 
            1.25: "1.25倍行距", 1.5: "1.5倍行距", 2.0: "双倍行距"
        }
    
    if not spacing_map:
        return f"{spacing}倍行距"
    
    closest_spacing = min(spacing_map.keys(), key=lambda x: abs(x - spacing))
    return spacing_map[closest_spacing]

def get_alignment_name(alignment_value, tpl=None):
    """段落对齐方式转换为中文描述"""
    alignment_map = {
        0: "左对齐", 1: "居中对齐", 2: "右对齐", 3: "两端对齐"
    }
    return alignment_map.get(alignment_value, f"未知对齐({alignment_value})")

def get_style_alignment(doc, style_id):
    """从文档的styles.xml中查找并返回指定样式的对齐方式"""
    try:
        styles = doc.styles.element
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        style_path = f".//w:style[@w:styleId='{style_id}']//w:jc"
        jc_element = styles.find(style_path, ns)
        if jc_element is not None:
            val = jc_element.get(ns['w'] + 'val')
            mapping = {'left': 0, 'center': 1, 'right': 2, 'both': 3, 'justify': 3}
            return mapping.get(val)
    except Exception:
        return None
    return None

def detect_paragraph_alignment(paragraph):
    """
    改进的段落对齐检测，按照明确的优先级处理直接格式和样式格式
    """
    # 优先级1: 检查直接格式
    direct_alignment = paragraph.paragraph_format.alignment
    if direct_alignment is not None:
        return int(direct_alignment)
    
    # 优先级2: 如果直接格式为None，检查样式格式
    if paragraph.style:
        try:
            style_alignment = paragraph.style.paragraph_format.alignment
            if style_alignment is not None:
                return int(style_alignment)
        except Exception:
            pass
    
    # 优先级3: 检查XML样式定义（处理复杂的样式继承）
    if paragraph.style and paragraph.style.style_id:
        try:
            xml_alignment = get_style_alignment(paragraph.part.document, paragraph.style.style_id)
            if xml_alignment is not None:
                return xml_alignment
        except Exception:
            pass
            
    # 默认值: 左对齐
    return WD_PARAGRAPH_ALIGNMENT.LEFT

def detect_paragraph_indent(paragraph):
    """检测段落缩进（返回pt值）"""
    try:
        fmt = paragraph.paragraph_format
        first_line_indent = fmt.first_line_indent.pt if fmt.first_line_indent else 0.0
        left_indent = fmt.left_indent.pt if fmt.left_indent else 0.0
        right_indent = fmt.right_indent.pt if fmt.right_indent else 0.0
        return first_line_indent, left_indent, right_indent
    except Exception:
        return 0.0, 0.0, 0.0

def pt_to_chars(pt_value, font_size_pt=12):
    """
    将pt值转换为字符数（用于首行缩进检测）
    假设1字符 ≈ font_size_pt pt
    """
    if pt_value == 0:
        return 0
    # 中文字符宽度约为字号大小，所以除以字号得到字符数
    return round(pt_value / font_size_pt, 1)
    
# ---------- 摘要检测逻辑 ----------
def check_abstract_structure(doc, tpl):
    """
    检查摘要结构（"摘 要"标题格式和内容长度）
    返回 {'ok': bool, 'messages': [], 'title_paragraph': paragraph, 'content_paragraphs': [paragraphs], 'content_text': str}
    """
    report = {'ok': True, 'messages': [], 'title_paragraph': None, 'content_paragraphs': [], 'content_text': ''}
    
    structure_rules = tpl.get('structure_rules', {})
    header_pattern = structure_rules.get('header_pattern', r'^\\s*摘\\s要\\s*$')
    content_start_line = structure_rules.get('content_start_line', 2)  # 标题后第几行开始是内容
    
    # 查找"摘 要"标题段落
    title_para = None
    title_idx = None
    
    # 使用模板中配置的标题正则（默认匹配“摘 要”，中间一个空格）
    title_pattern = header_pattern if header_pattern else r'^\s*摘\s要\s*$'
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if re.match(title_pattern, text):
            title_para = paragraph
            title_idx = idx
            break
    
    if not title_para:
        report['ok'] = False
        error_msg = tpl.get('messages', {}).get('structure_header_error')
        if error_msg:
            report['messages'].append(error_msg)
        else:
            report['messages'].append("未找到'摘 要'标题段落")
        return report
    
    report['title_paragraph'] = title_para
    # save paragraph index for downstream consumers
    report['title_paragraph_index'] = title_idx
    # record paragraph index for robust downstream consumption
    try:
        report['title_paragraph_index'] = title_idx
    except Exception:
        report['title_paragraph_index'] = None
    
    # 检查标题格式（"摘 要"中间应有一个空格）
    title_text = title_para.text.strip()
    if not re.match(r'^\s*摘\s要\s*$', title_text):
        # 检查空格数量
        match = re.search(r'摘(\s*)要', title_text)
        if match:
            space_count = len(match.group(1))
            if space_count != 1:
                report['ok'] = False
                report['messages'].append(f"'摘 要'中间应有1个空格（当前有{space_count}个空格）")
        else:
            report['ok'] = False
            report['messages'].append("'摘 要'标题格式错误，应为'摘 要'（中间一个空格）")
    else:
        ok_msg = tpl.get('messages', {}).get('structure_header_ok')
        if ok_msg:
            report['messages'].append(ok_msg)
    
    # 查找摘要正文段落（标题后的第content_start_line行开始）
    content_paragraphs = []
    content_start_idx = title_idx + content_start_line
    
    # 从content_start_idx开始查找所有非空段落，直到遇到下一个标题（如"关键词"）
    stop_keywords = ['关键词', '关键字', 'Keywords', 'Key words']
    
    for idx in range(content_start_idx, len(doc.paragraphs)):
        para = doc.paragraphs[idx]
        text = para.text.strip()
        
        if not text:
            continue
        
        # 遇到停止关键词，停止查找
        if any(keyword in text for keyword in stop_keywords):
            break
        
        # 判断是否包含中文字符（摘要内容应该包含中文）
        if re.search(r'[\u4e00-\u9fff]', text):
            content_paragraphs.append(para)
    
    if not content_paragraphs:
        report['ok'] = False
        report['messages'].append("未找到摘要正文内容")
        return report
    
    report['content_paragraphs'] = content_paragraphs
    # also store numeric indices for content paragraphs to make locating deterministic
    try:
        indices = []
        for para in content_paragraphs:
            for i, p in enumerate(doc.paragraphs):
                if p is para:
                    indices.append(i)
                    break
        report['content_paragraphs_indices'] = indices
    except Exception:
        report['content_paragraphs_indices'] = []
    # 合并所有正文段落的内容
    content_text = ' '.join([para.text.strip() for para in content_paragraphs])
    report['content_text'] = content_text
    # record content paragraph indexes
    try:
        content_indexes = []
        for p in content_paragraphs:
            for i, para in enumerate(doc.paragraphs):
                if para is p:
                    content_indexes.append(i)
                    break
        report['content_paragraph_indexes'] = content_indexes
    except Exception:
        report['content_paragraph_indexes'] = []
    
    # 检查内容长度
    content_length = len(content_text)
    min_length = structure_rules.get('min_content_length', 300)
    max_length = structure_rules.get('max_content_length', 2000)
    
    if content_length < min_length:
        report['ok'] = False
        msg_tpl = tpl.get('messages', {}).get('structure_length_short')
        if msg_tpl:
            try:
                report['messages'].append(msg_tpl.format(min=min_length))
            except:
                report['messages'].append(msg_tpl)
    elif content_length > max_length:
        report['ok'] = False
        msg_tpl = tpl.get('messages', {}).get('structure_length_long')
        if msg_tpl:
            try:
                report['messages'].append(msg_tpl.format(max=max_length))
            except:
                report['messages'].append(msg_tpl)
    else:
        ok_msg = tpl.get('messages', {}).get('structure_length_ok')
        if ok_msg:
            report['messages'].append(ok_msg)
    
    return report

def check_abstract_title_format(paragraph, tpl):
    """
    检查摘要标题格式（"摘 要"）
    返回 {'ok': bool, 'messages': []}
    """
    report = {'ok': True, 'messages': []}
    
    if not paragraph or not paragraph.runs:
        report['ok'] = False
        report['messages'].append("摘要标题段落没有文本内容")
        return report
    
    format_rules = tpl.get('format_rules', {}).get('abstract', {})
    
    # 检测第一个非空run的格式
    main_run = None
    for run in paragraph.runs:
        if run.text.strip():
            main_run = run
            break
    
    if not main_run:
        report['ok'] = False
        report['messages'].append("摘要标题段落没有有效文本")
        return report
    
    # 检测实际格式
    actual_size_pt, actual_font_ascii, actual_font_eastasia, actual_bold, actual_italic, actual_line_spacing = detect_font_for_run(main_run, paragraph)
    
    issues = []
    
    # 字体大小检查
    if not should_skip_check('font_size') and 'font_size_pt' in format_rules:
        expected_size_pt = float(format_rules['font_size_pt'])
        actual_size_name = get_font_size(actual_size_pt, tpl)
        expected_size_name = get_font_size(expected_size_pt, tpl)
        print(f"标题字体大小: {actual_size_name}（{actual_size_pt}pt）(期望: {expected_size_name}（{expected_size_pt}pt）)")
        if abs(actual_size_pt - expected_size_pt) > 0.5:
            issues.append(f"标题字体大小应为{expected_size_name}（{expected_size_pt}pt），实际为{actual_size_name}（{actual_size_pt}pt）")
    
    # 中文字体检查（标题应为黑体）
    if not should_skip_check('font_name') and 'font_name' in format_rules:
        expected_font_name = str(format_rules['font_name'])
        print(f"标题中文字体: {actual_font_eastasia} (期望: {expected_font_name})")
        if expected_font_name.lower() not in actual_font_eastasia.lower():
            issues.append(f"标题中文字体应为{expected_font_name}，实际为{actual_font_eastasia}")
    
    # 加粗检查
    if not should_skip_check('bold') and 'bold' in format_rules:
        expected_bold = bool(format_rules['bold'])
        print(f"标题加粗: {'是' if actual_bold else '否'} (期望: {'是' if expected_bold else '否'})")
        if actual_bold != expected_bold:
            bold_status = "加粗" if expected_bold else "不加粗"
            actual_status = "加粗" if actual_bold else "不加粗"
            issues.append(f"标题字体应为{bold_status}，实际为{actual_status}")
    
    # 行间距检查
    if not should_skip_check('spacing') and 'line_spacing' in format_rules:
        expected_line_spacing = float(format_rules['line_spacing'])
        actual_spacing_name = get_line_spacing_name(actual_line_spacing, tpl)
        expected_spacing_name = get_line_spacing_name(expected_line_spacing, tpl)
        print(f"标题行间距: {actual_spacing_name}（{actual_line_spacing}倍）(期望: {expected_spacing_name}（{expected_line_spacing}倍）)")
        if abs(actual_line_spacing - expected_line_spacing) > 0.1:
            issues.append(f"标题行间距应为{expected_spacing_name}（{expected_line_spacing}倍），实际为{actual_spacing_name}（{actual_line_spacing}倍）")
    
    # 段落对齐检查
    if 'alignment' in format_rules:
        expected_alignment_str = str(format_rules['alignment'])
        alignment_map = {"left": 0, "center": 1, "right": 2, "justify": 3}
        expected_alignment = alignment_map.get(expected_alignment_str, 0)
        
        actual_alignment = detect_paragraph_alignment(paragraph)
        actual_alignment_name = get_alignment_name(actual_alignment, tpl)
        expected_alignment_name = get_alignment_name(expected_alignment, tpl)
        print(f"标题对齐: {actual_alignment_name} (期望: {expected_alignment_name})")
        if actual_alignment != expected_alignment:
            issues.append(f"标题应为{expected_alignment_name}，实际为{actual_alignment_name}")
    
    # 段落缩进检查（标题应无缩进）
    if 'first_line_indent' in format_rules or 'left_indent' in format_rules:
        first_line_indent, left_indent, right_indent = detect_paragraph_indent(paragraph)
        
        if 'first_line_indent' in format_rules:
            expected_first_indent = float(format_rules['first_line_indent'])
            print(f"标题首行缩进: {first_line_indent:.1f}pt (期望: {expected_first_indent}pt)")
            if abs(first_line_indent - expected_first_indent) > 1.0:
                issues.append(f"标题首行缩进应为{expected_first_indent}pt，实际为{first_line_indent:.1f}pt")
        
        if 'left_indent' in format_rules:
            expected_left_indent = float(format_rules['left_indent'])
            print(f"标题左缩进: {left_indent:.1f}pt (期望: {expected_left_indent}pt)")
            if abs(left_indent - expected_left_indent) > 1.0:
                issues.append(f"标题左缩进应为{expected_left_indent}pt，实际为{left_indent:.1f}pt")
    
    print(f"标题格式检查发现 {len(issues)} 个问题")
    print("---")
    
    if issues:
        report['ok'] = False
        header = tpl.get('messages', {}).get('format_abstract_issue_header')
        if header:
            report['messages'].append(header)
        report['messages'].extend([f"  - {i}" for i in issues])
    else:
        ok_msg = tpl.get('messages', {}).get('format_abstract_ok')
        if ok_msg:
            report['messages'].append(ok_msg)
    
    return report

def check_abstract_content_format(content_paragraphs, tpl):
    """
    检查摘要正文格式
    返回 {'ok': bool, 'messages': []}
    """
    report = {'ok': True, 'messages': []}
    
    if not content_paragraphs:
        report['ok'] = False
        report['messages'].append("摘要正文段落为空")
        return report
    
    format_rules = tpl.get('format_rules', {}).get('content', {})
    issues = []
    
    # 检查每个正文段落
    for para_idx, paragraph in enumerate(content_paragraphs, 1):
        if not paragraph or not paragraph.runs:
            continue
        
        # 检测第一个非空run的格式
        main_run = None
        for run in paragraph.runs:
            if run.text.strip():
                main_run = run
                break
        
        if not main_run:
            continue
        
        # 检测实际格式
        actual_size_pt, actual_font_ascii, actual_font_eastasia, actual_bold, actual_italic, actual_line_spacing = detect_font_for_run(main_run, paragraph)
        
        # 字体大小检查
        if not should_skip_check('font_size') and 'font_size_pt' in format_rules:
            expected_size_pt = float(format_rules['font_size_pt'])
            actual_size_name = get_font_size(actual_size_pt, tpl)
            expected_size_name = get_font_size(expected_size_pt, tpl)
            if abs(actual_size_pt - expected_size_pt) > 0.5:
                msg = f"第{para_idx}段正文字体大小应为{expected_size_name}（{expected_size_pt}pt），实际为{actual_size_name}（{actual_size_pt}pt）"
                if msg not in issues:
                    issues.append(msg)
        
        # 检查run中的中文字符
        if re.search(r'[\u4e00-\u9fff]', main_run.text):
            expected_chinese_font = format_rules.get('chinese_font', '宋体')
            if expected_chinese_font.lower() not in actual_font_eastasia.lower():
                msg = f"第{para_idx}段正文中文字体应为{expected_chinese_font}，实际为{actual_font_eastasia}"
                if msg not in issues:
                    issues.append(msg)
        
        # 检查run中的英文字符
        if re.search(r'[a-zA-Z]', main_run.text):
            expected_english_font = format_rules.get('english_font', 'Times New Roman')
            if expected_english_font.lower() not in actual_font_ascii.lower():
                msg = f"第{para_idx}段正文英文字体应为{expected_english_font}，实际为{actual_font_ascii}"
                if msg not in issues:
                    issues.append(msg)
        
        # 加粗检查（正文不应加粗）
        if not should_skip_check('bold') and 'bold' in format_rules:
            expected_bold = bool(format_rules['bold'])
            if actual_bold != expected_bold:
                bold_status = "加粗" if expected_bold else "不加粗"
                actual_status = "加粗" if actual_bold else "不加粗"
                msg = f"第{para_idx}段正文应为{bold_status}，实际为{actual_status}"
                if msg not in issues:
                    issues.append(msg)
        
        # 行间距检查
        if not should_skip_check('spacing') and 'line_spacing' in format_rules:
            expected_line_spacing = float(format_rules['line_spacing'])
            if abs(actual_line_spacing - expected_line_spacing) > 0.1:
                actual_spacing_name = get_line_spacing_name(actual_line_spacing, tpl)
                expected_spacing_name = get_line_spacing_name(expected_line_spacing, tpl)
                msg = f"第{para_idx}段正文行间距应为{expected_spacing_name}（{expected_line_spacing}倍），实际为{actual_spacing_name}（{actual_line_spacing}倍）"
                if msg not in issues:
                    issues.append(msg)
        
        # 段落对齐检查
        if 'alignment' in format_rules:
            expected_alignment_str = str(format_rules['alignment'])
            alignment_map = {"left": 0, "center": 1, "right": 2, "justify": 3}
            expected_alignment = alignment_map.get(expected_alignment_str, 0)
            
            actual_alignment = detect_paragraph_alignment(paragraph)
            if actual_alignment != expected_alignment:
                actual_alignment_name = get_alignment_name(actual_alignment, tpl)
                expected_alignment_name = get_alignment_name(expected_alignment, tpl)
                msg = f"第{para_idx}段正文应为{expected_alignment_name}，实际为{actual_alignment_name}"
                if msg not in issues:
                    issues.append(msg)
        
        # 首行缩进检查（转换为字符数）
        if 'first_line_indent' in format_rules:
            first_line_indent, left_indent, right_indent = detect_paragraph_indent(paragraph)
            expected_indent_chars = float(format_rules['first_line_indent'])  # 期望的字符数
            expected_indent_pt = expected_indent_chars * actual_size_pt  # 转换为pt
            actual_indent_chars = pt_to_chars(first_line_indent, actual_size_pt)
            
            if abs(first_line_indent - expected_indent_pt) > actual_size_pt * 0.5:  # 容差为0.5字符
                msg = f"第{para_idx}段正文首行缩进应为{expected_indent_chars}字符，实际约为{actual_indent_chars}字符"
                if msg not in issues:
                    issues.append(msg)
    
    print(f"正文格式检查发现 {len(issues)} 个问题")
    print("---")
    
    if issues:
        report['ok'] = False
        header = tpl.get('messages', {}).get('format_abstract_issue_header')
        if header:
            report['messages'].append(header)
        report['messages'].extend([f"  - {i}" for i in issues])
    else:
        ok_msg = tpl.get('messages', {}).get('format_content_ok')
        if ok_msg:
            report['messages'].append(ok_msg)
    
    return report

def check_abstract_with_template(doc_path, template_identifier, skip_checks=None):
    """
    主检查函数：检查中文摘要格式
    参数:
        doc_path: 文档路径
        template_identifier: 模板标识符
        skip_checks: 要跳过的检测项列表，如 ['font_size', 'bold']
    """
    # 设置全局跳过检测项配置
    global _skip_checks_config
    _skip_checks_config = skip_checks or []
    
    tpl = load_template(template_identifier)
    doc = Document(doc_path)
    
    # 执行各项检查
    structure_report = check_abstract_structure(doc, tpl)
    
    title_format_report = {'ok': True, 'messages': []}
    content_format_report = {'ok': True, 'messages': []}
    
    if structure_report.get('title_paragraph'):
        title_format_report = check_abstract_title_format(structure_report['title_paragraph'], tpl)
    
    if structure_report.get('content_paragraphs'):
        content_format_report = check_abstract_content_format(structure_report['content_paragraphs'], tpl)
    
    # 组装报告
    report = {
        'structure': structure_report,
        'title_format': title_format_report,
        'content_format': content_format_report,
        'summary': []
    }
    
    # 生成总结
    all_ok = (structure_report['ok'] and title_format_report['ok'] and content_format_report['ok'])
    summary_tpl = tpl.get('messages', {}).get('summary_overall')
    if summary_tpl:
        try:
            report['summary'].append(summary_tpl.format(ok="通过" if all_ok else "失败"))
        except Exception:
            report['summary'].append(str(summary_tpl))
    
    return report

# ---------- 报告输出 ----------
def print_abstract_report(report):
    """打印中文摘要检查报告"""
    print("=== 中文摘要检查报告 ===")
    
    sections = [
        ('structure', '结构检测'),
        ('title_format', '标题格式'),
        ('content_format', '正文格式')
    ]
    
    for sec_key, sec_name in sections:
        info = report.get(sec_key, {})
        print(f"--- {sec_name} ---")
        print(" 状态:", "✓ 通过" if info.get('ok', False) else "✗ 失败")
        for m in info.get('messages', []):
            print("  -", m)
    
    print("--- 总结 ---")
    for s in report.get('summary', []):
        print(" ", s)

def print_help():
    print("使用方法:")
    print("  python Abstract_detect.py check_cn <paper.docx> <cn_template>")
    print("  python Abstract_detect.py check_en <paper.docx> <en_template>")
    print("  python Abstract_detect.py check_bi <paper.docx> <cn_template> <en_template>")

if __name__ == '__main__':
    cmd = sys.argv[1] if len(sys.argv) > 1 else None
    if cmd == 'check_cn' and len(sys.argv) == 4:
        paper_path = sys.argv[2]
        tpl_id = sys.argv[3]
        if not os.path.isfile(paper_path):
            print(f"论文文件不存在: {paper_path}")
            sys.exit(1)
        try:
            print("=== 开始中文摘要格式检查 ===")
            report = check_abstract_with_template(paper_path, tpl_id)
        except Exception as e:
            print("检查时出错:", e)
            import traceback
            traceback.print_exc()
            sys.exit(1)
        print_abstract_report(report)
    elif cmd == 'check_en' and len(sys.argv) == 4:
        paper_path = sys.argv[2]
        tpl_id = sys.argv[3]
        if not os.path.isfile(paper_path):
            print(f"论文文件不存在: {paper_path}")
            sys.exit(1)
        try:
            print("=== 开始英文摘要格式检查 ===")
            report = check_english_abstract_with_template(paper_path, tpl_id)
        except Exception as e:
            print("检查时出错:", e)
            import traceback
            traceback.print_exc()
            sys.exit(1)
        print("=== 英文摘要检查报告 ===")
        print(report)
    elif cmd == 'check_bi' and len(sys.argv) == 5:
        paper_path = sys.argv[2]
        cn_tpl = sys.argv[3]
        en_tpl = sys.argv[4]
        if not os.path.isfile(paper_path):
            print(f"论文文件不存在: {paper_path}")
            sys.exit(1)
        try:
            print("=== 开始中英文摘要格式检查 ===")
            report = check_bilingual_abstracts(paper_path, cn_tpl, en_tpl)
        except Exception as e:
            print("检查时出错:", e)
            import traceback
            traceback.print_exc()
            sys.exit(1)
        # 打印简要报告
        print("=== 中英文摘要检查报告 ===")
        print("  - 总结:", "; ".join(report.get('summary', [])))
        print("\n[中文摘要]")
        print(report.get('chinese'))
        print("\n[英文摘要]")
        print(report.get('english'))
    else:
        print_help()
        sys.exit(0)

