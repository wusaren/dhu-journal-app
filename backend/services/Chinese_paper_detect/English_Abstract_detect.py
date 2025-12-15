#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
英文摘要检测器
依据模板 Chinese_paper_detect_templates/English_Abstract.json
"""

import os
import sys
import json
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

# 便于独立运行
if __name__ == "__main__" and __package__ is None:
    file = Path(__file__).resolve()
    parent, root = file.parent, file.parents[1]
    sys.path.append(str(root))
    try:
        sys.path.remove(str(parent))
    except ValueError:
        pass

_skip_checks_config = []


def should_skip_check(check_name):
    global _skip_checks_config
    return check_name in (_skip_checks_config or [])


# ---------- 模板加载 ----------
def resolve_template_path(identifier):
    if os.path.isfile(identifier):
        return identifier
    current_dir = Path(__file__).parent
    candidates = [
        os.path.join("templates", identifier + ".json"),
        os.path.join("Chinese_paper_detect_templates", identifier + ".json"),
        os.path.join("services", "Chinese_paper_detect_templates", identifier + ".json"),
        str(current_dir.parent / "Chinese_paper_detect_templates" / (identifier + ".json")),
        str(current_dir / ".." / "Chinese_paper_detect_templates" / (identifier + ".json")),
    ]
    for candidate in candidates:
        abs_path = os.path.abspath(candidate)
        if os.path.isfile(abs_path):
            return abs_path
        if os.path.isfile(candidate):
            return candidate
    raise FileNotFoundError(f"Template not found: '{identifier}' (tried {candidates})")


def load_template(identifier):
    tpl_path = resolve_template_path(identifier)
    with open(tpl_path, "r", encoding="utf-8") as f:
        return json.load(f)


# ---------- 基础检测工具 ----------
def detect_font_for_run(run, paragraph=None):
    font_size = None
    font_name_ascii = None
    is_bold = None
    is_italic = False

    if not run:
        return 12.0, "Times New Roman", False, False, 1.0

    # 字号：直接格式
    try:
        if run.font and run.font.size and hasattr(run.font.size, "pt"):
            font_size = float(run.font.size.pt)

        # run.style
        if font_size is None and getattr(run, "style", None) and getattr(run.style, "font", None):
            if run.style.font.size and hasattr(run.style.font.size, "pt"):
                font_size = float(run.style.font.size.pt)

        # paragraph.style
        if font_size is None and paragraph and paragraph.style and getattr(paragraph.style, "font", None):
            if paragraph.style.font.size and hasattr(paragraph.style.font.size, "pt"):
                font_size = float(paragraph.style.font.size.pt)

        # 段落样式 XML
        if font_size is None and paragraph and paragraph.style and hasattr(paragraph.style, "element"):
            sz_nodes = paragraph.style.element.xpath(".//w:sz")
            if sz_nodes and sz_nodes[0].get(qn("w:val")):
                font_size = float(sz_nodes[0].get(qn("w:val"))) / 2.0

        # run 的 rPr
        if font_size is None and hasattr(run._element, "rPr"):
            sz_nodes = run._element.xpath(".//w:sz")
            if sz_nodes and sz_nodes[0].get(qn("w:val")):
                font_size = float(sz_nodes[0].get(qn("w:val"))) / 2.0
    except Exception:
        pass
    font_size = font_size if font_size is not None else 12.0

    # 字体（英文）
    try:
        if run.font and run.font.name:
            font_name_ascii = run.font.name
        if hasattr(run._element, "rPr"):
            rpr = run._element.rPr
            if rpr is not None:
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is not None:
                    xml_ascii = rfonts.get(qn("w:ascii"))
                    xml_hansi = rfonts.get(qn("w:hAnsi"))
                    if xml_ascii:
                        font_name_ascii = xml_ascii
                    elif xml_hansi and font_name_ascii is None:
                        font_name_ascii = xml_hansi
    except Exception:
        pass
    font_name_ascii = font_name_ascii if font_name_ascii else "Times New Roman"

    # 加粗/斜体
    try:
        if run.font:
            if run.font.bold is not None:
                is_bold = run.font.bold
            if run.font.italic is not None:
                is_italic = run.font.italic

        if is_bold is None and getattr(run, "style", None) and getattr(run.style, "font", None):
            if run.style.font.bold is not None:
                is_bold = run.style.font.bold

        if is_bold is None and paragraph and paragraph.style and getattr(paragraph.style, "font", None):
            if paragraph.style.font.bold is not None:
                is_bold = paragraph.style.font.bold

        if is_bold is None and hasattr(run._element, "rPr"):
            b_nodes = run._element.xpath(".//w:b")
            if b_nodes:
                is_bold = True

        is_bold = bool(is_bold) if is_bold is not None else False
        is_italic = bool(is_italic)
    except Exception:
        pass

    # 行间距
    line_spacing = 1.0
    try:
        if paragraph and paragraph.paragraph_format.line_spacing:
            line_spacing = float(paragraph.paragraph_format.line_spacing)
        elif paragraph and paragraph.style and paragraph.style.paragraph_format.line_spacing:
            line_spacing = float(paragraph.style.paragraph_format.line_spacing)
    except Exception:
        pass

    return font_size, font_name_ascii, is_bold, is_italic, line_spacing


def get_font_size(pt_size):
    size_map = {9: "小五", 10.5: "五号", 11: "小四(11pt)", 12: "小四", 14: "四号", 16: "三号"}
    if not size_map:
        return f"{pt_size}pt"
    closest_size = min(size_map.keys(), key=lambda x: abs(x - pt_size))
    return size_map[closest_size]


def get_line_spacing_name(spacing):
    spacing_map = {1.0: "单倍行距", 1.15: "1.15倍行距", 1.25: "1.25倍行距", 1.5: "1.5倍行距", 2.0: "双倍行距"}
    closest_spacing = min(spacing_map.keys(), key=lambda x: abs(x - spacing))
    return spacing_map[closest_spacing]


def get_alignment_name(alignment_value):
    alignment_map = {0: "左对齐", 1: "居中对齐", 2: "右对齐", 3: "两端对齐"}
    return alignment_map.get(alignment_value, f"未知对齐({alignment_value})")


def get_style_alignment(doc, style_id):
    try:
        styles = doc.styles.element
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        style_path = f".//w:style[@w:styleId='{style_id}']//w:jc"
        jc_element = styles.find(style_path, ns)
        if jc_element is not None:
            val = jc_element.get(ns["w"] + "val")
            mapping = {"left": 0, "center": 1, "right": 2, "both": 3, "justify": 3}
            return mapping.get(val)
    except Exception:
        return None
    return None


def detect_paragraph_alignment(paragraph):
    direct_alignment = paragraph.paragraph_format.alignment
    if direct_alignment is not None:
        return int(direct_alignment)
    if paragraph.style:
        try:
            style_alignment = paragraph.style.paragraph_format.alignment
            if style_alignment is not None:
                return int(style_alignment)
        except Exception:
            pass
    if paragraph.style and paragraph.style.style_id:
        try:
            xml_alignment = get_style_alignment(paragraph.part.document, paragraph.style.style_id)
            if xml_alignment is not None:
                return xml_alignment
        except Exception:
            pass
    return WD_PARAGRAPH_ALIGNMENT.LEFT


def detect_paragraph_indent(paragraph):
    try:
        fmt = paragraph.paragraph_format
        first_line_indent = fmt.first_line_indent.pt if fmt.first_line_indent else 0.0
        left_indent = fmt.left_indent.pt if fmt.left_indent else 0.0
        right_indent = fmt.right_indent.pt if fmt.right_indent else 0.0
        return first_line_indent, left_indent, right_indent
    except Exception:
        return 0.0, 0.0, 0.0


def pt_to_chars(pt_value, font_size_pt=12):
    if pt_value == 0:
        return 0
    return round(pt_value / font_size_pt, 1)


# ---------- 英文摘要检测 ----------
def check_english_abstract(doc, tpl):
    # 分段报告
    structure_report = {"ok": True, "messages": []}
    title_report = {"ok": True, "messages": []}
    content_report = {"ok": True, "messages": []}
    report = {"ok": True, "messages": [], "title_paragraph": None, "content_paragraphs": []}
    sr = tpl.get("structure_rules", {})
    header_pattern = sr.get("header_pattern", r"^\s*ABSTRACT\s*$")
    blank_after_title = sr.get("blank_lines_after_title", 1)
    min_len = sr.get("min_content_length", 50)
    max_len = sr.get("max_content_length", 2000)

    # 标题定位
    title_idx = None
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if re.match(header_pattern, text, re.IGNORECASE):
            title_idx = idx
            report["title_paragraph"] = para
            break
    if title_idx is None:
        structure_report["ok"] = False
        msg = tpl.get("messages", {}).get("structure_header_error", "未找到 ABSTRACT 标题")
        structure_report["messages"].append(msg)
        report["ok"] = False
        report["messages"].append(msg)
        return {
            "ok": False,
            "structure": structure_report,
            "title_format": title_report,
            "content_format": content_report,
            "summary": []
        }
    ok_msg = tpl.get("messages", {}).get("structure_header_ok")
    if ok_msg:
        structure_report["messages"].append(ok_msg)

    # 标题与正文间空行
    blank_count = 0
    cursor = title_idx + 1
    while cursor < len(doc.paragraphs) and not doc.paragraphs[cursor].text.strip():
        blank_count += 1
        cursor += 1
    if blank_count < blank_after_title:
        structure_report["ok"] = False
        report["ok"] = False
        msg = tpl.get("messages", {}).get("structure_blank_after_title_error", "标题与正文之间应空一行")
        structure_report["messages"].append(msg)

    # 正文段落收集
    content_paras = []
    consecutive_blanks = 0
    for idx in range(cursor, len(doc.paragraphs)):
        para = doc.paragraphs[idx]
        text = para.text.strip()
        if not text:
            consecutive_blanks += 1
            if consecutive_blanks >= sr.get("blank_lines_between_sections", 2):
                break
            continue
        consecutive_blanks = 0
        if re.search(r"Supervised by", text, re.IGNORECASE):
            break
        content_paras.append(para)

    if not content_paras:
        structure_report["ok"] = False
        report["ok"] = False
        structure_report["messages"].append("未找到英文摘要正文内容")
        return {
            "ok": False,
            "structure": structure_report,
            "title_format": title_report,
            "content_format": content_report,
            "summary": []
        }

    report["content_paragraphs"] = content_paras

    # 长度检查
    content_text = " ".join([p.text.strip() for p in content_paras])
    if len(content_text) < min_len:
        structure_report["ok"] = False
        report["ok"] = False
        msg = tpl.get("messages", {}).get("structure_length_short", f"英文摘要过短（少于 {min_len} 字符）")
        try:
            structure_report["messages"].append(msg.format(min=min_len))
        except Exception:
            structure_report["messages"].append(msg)
    elif len(content_text) > max_len:
        structure_report["ok"] = False
        report["ok"] = False
        msg = tpl.get("messages", {}).get("structure_length_long", f"英文摘要过长（超过 {max_len} 字符）")
        try:
            structure_report["messages"].append(msg.format(max=max_len))
        except Exception:
            structure_report["messages"].append(msg)
    else:
        ok_msg = tpl.get("messages", {}).get("structure_length_ok")
        if ok_msg:
            structure_report["messages"].append(ok_msg)

    # 标题格式
    title_rules = tpl.get("format_rules", {}).get("title", {})
    if report["title_paragraph"]:
        runs = report["title_paragraph"].runs
        main_run = next((r for r in runs if r.text.strip()), None)
        if main_run:
            size_pt, font_ascii, bold, italic, line_spacing = detect_font_for_run(main_run, report["title_paragraph"])
            issues = []
            if "font_size_pt" in title_rules and abs(size_pt - float(title_rules["font_size_pt"])) > 0.5:
                issues.append(f"标题字号应为{title_rules['font_size_pt']}pt，当前{size_pt}pt")
            if "font_name" in title_rules:
                expected_font = str(title_rules["font_name"])
                if expected_font.lower() not in font_ascii.lower():
                    issues.append(f"标题字体应为{expected_font}，当前{font_ascii}")
            if "bold" in title_rules and bool(bold) != bool(title_rules["bold"]):
                issues.append("标题应加粗" if title_rules["bold"] else "标题不应加粗")
            if "alignment" in title_rules:
                alignment_map = {"left": 0, "center": 1, "right": 2, "justify": 3}
                expected_align = alignment_map.get(str(title_rules["alignment"]), 1)
                actual_align = detect_paragraph_alignment(report["title_paragraph"])
                if actual_align != expected_align:
                    issues.append(f"标题应居中对齐，当前对齐值 {actual_align}")
            if "space_before" in title_rules or "space_after" in title_rules:
                space_before = report["title_paragraph"].paragraph_format.space_before
                space_after = report["title_paragraph"].paragraph_format.space_after
                def pt(v): return v.pt if v else 0
                if "space_before" in title_rules and abs(pt(space_before) - float(title_rules["space_before"])) > 0.1:
                    issues.append(f"标题段前应为{title_rules['space_before']}，当前{pt(space_before)}")
                if "space_after" in title_rules and abs(pt(space_after) - float(title_rules["space_after"])) > 0.1:
                    issues.append(f"标题段后应为{title_rules['space_after']}，当前{pt(space_after)}")
            if "line_spacing" in title_rules:
                if line_spacing and abs(float(line_spacing) - float(title_rules["line_spacing"])) > 0.1:
                    issues.append(f"标题行距应为{title_rules['line_spacing']}倍，当前{line_spacing}")

            if issues:
                report["ok"] = False
                title_report["ok"] = False
                header = tpl.get("messages", {}).get("format_issue_header", "英文摘要格式问题：")
                title_report["messages"].append(header)
                title_report["messages"].extend([f"  - {i}" for i in issues])
            else:
                ok_msg = tpl.get("messages", {}).get("format_title_ok")
                if ok_msg:
                    title_report["messages"].append(ok_msg)

    # 正文格式
    content_rules = tpl.get("format_rules", {}).get("content", {})
    content_issues = []
    for idx, para in enumerate(content_paras, 1):
        main_run = next((r for r in para.runs if r.text.strip()), None)
        if not main_run:
            continue
        size_pt, font_ascii, bold, italic, line_spacing = detect_font_for_run(main_run, para)
        if "font_size_pt" in content_rules and abs(size_pt - float(content_rules["font_size_pt"])) > 0.5:
            content_issues.append(f"第{idx}段字号应为{content_rules['font_size_pt']}pt，当前{size_pt}pt")
        expected_font = content_rules.get("english_font", "Times New Roman")
        if expected_font.lower() not in font_ascii.lower():
            content_issues.append(f"第{idx}段字体应为{expected_font}，当前{font_ascii}")
        if "bold" in content_rules and bool(bold) != bool(content_rules["bold"]):
            content_issues.append(f"第{idx}段不应加粗")
        if "line_spacing" in content_rules and line_spacing and abs(float(line_spacing) - float(content_rules["line_spacing"])) > 0.1:
            content_issues.append(f"第{idx}段行距应为{content_rules['line_spacing']}倍，当前{line_spacing}")
        if "alignment" in content_rules:
            alignment_map = {"left": 0, "center": 1, "right": 2, "justify": 3}
            expected_align = alignment_map.get(str(content_rules["alignment"]), 0)
            actual_align = detect_paragraph_alignment(para)
            if actual_align != expected_align:
                content_issues.append("第{}段应左对齐".format(idx))
        if "first_line_indent_chars" in content_rules:
            first_line_indent, _, _ = detect_paragraph_indent(para)
            expected_chars = float(content_rules["first_line_indent_chars"])
            expected_pt = expected_chars * size_pt
            actual_chars = pt_to_chars(first_line_indent, size_pt)
            if abs(first_line_indent - expected_pt) > size_pt * 0.5:
                content_issues.append(f"第{idx}段首行缩进应为{expected_chars}字符，当前约{actual_chars}字符")

    if content_issues:
        report["ok"] = False
        content_report["ok"] = False
        header = tpl.get("messages", {}).get("format_issue_header", "英文摘要格式问题：")
        content_report["messages"].append(header)
        content_report["messages"].extend([f"  - {i}" for i in content_issues])
    else:
        ok_msg = tpl.get("messages", {}).get("format_content_ok")
        if ok_msg:
            content_report["messages"].append(ok_msg)

    # 汇总
    overall_ok = structure_report["ok"] and title_report["ok"] and content_report["ok"]
    report = {
        "ok": overall_ok,
        "structure": structure_report,
        "title_format": title_report,
        "content_format": content_report,
        "summary": []
    }
    summary_tpl = tpl.get("messages", {}).get("summary_overall")
    if summary_tpl:
        try:
            report["summary"].append(summary_tpl.format(ok="通过" if overall_ok else "失败"))
        except Exception:
            report["summary"].append(str(summary_tpl))
    return report


def check_english_abstract_with_template(doc_path, template_identifier, skip_checks=None):
    global _skip_checks_config
    _skip_checks_config = skip_checks or []
    tpl = load_template(template_identifier)
    doc = Document(doc_path)
    return check_english_abstract(doc, tpl)


def print_en_report(report):
    print("=== 英文摘要检查报告 ===")
    print(" 状态:", "✓ 通过" if report.get("ok") else "✗ 失败")
    for sec_key, sec_name in [
        ("structure", "Structure"),
        ("title_format", "Title Format"),
        ("content_format", "Content Format"),
    ]:
        info = report.get(sec_key, {})
        if not info:
            continue
        print(f"  [{sec_name}] {'✓ 通过' if info.get('ok') else '✗ 失败'}")
        for m in info.get("messages", []):
            print("    •", m)
    if "summary" in report:
        print("【总结】")
        for s in report["summary"]:
            print(" ", s)


def print_help():
    print("Usage:")
    print("  python English_Abstract_detect.py check <paper.docx> <template.json_or_name>")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print_help()
        sys.exit(0)
    cmd = sys.argv[1]
    if cmd == "check":
        paper_path = sys.argv[2]
        tpl_id = sys.argv[3]
        if not os.path.isfile(paper_path):
            print(f"论文文件不存在: {paper_path}")
            sys.exit(1)
        try:
            print("=== 开始英文摘要格式检查 ===")
            report = check_english_abstract_with_template(paper_path, tpl_id)
        except Exception as e:
            print("检查时出错:", e)
            import traceback
            traceback.print_exc()
            sys.exit(1)
        print_en_report(report)
    else:
        print_help()
        sys.exit(0)

