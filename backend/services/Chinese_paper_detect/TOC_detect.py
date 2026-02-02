#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
目录/图录/表录格式检测器
检测目录、图录、表录的格式是否符合要求
"""

import os
import sys
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn

# 添加项目根目录到 sys.path 以支持独立运行
if __name__ == "__main__" and __package__ is None:
    file = Path(__file__).resolve()
    parent, root = file.parent, file.parents[1]
    sys.path.append(str(root))
    try:
        sys.path.remove(str(parent))
    except ValueError:
        pass

def should_skip_check(check_name):
    """
    判断是否应该跳过某个检测项
    参数:
        check_name: 检测项名称 (font_size, bold, italic, alignment, spacing, indent)
    返回:
        bool: True表示跳过该检测项，False表示执行该检测项
    """
    global _skip_checks_config
    if _skip_checks_config is None:
        return False
    return check_name in _skip_checks_config

# 全局变量，用于存储当前模块的跳过检测项配置
_skip_checks_config = []

"""
=== 论文格式检测系统 - 目录/图录/表录检测器 ===

【目录/图录/表录检测 (TOC/Figure List/Table List Detection)】

1. 【结构检测】
   - 找到"目录"、"图录"、"表录"标题
   - 检查标题与条目之间是否空两行

2. 【格式检测】
   - 标题格式：黑体3号（16pt），段前0.7厘米，段后0
"""

# ---------- 模板加载 ----------
def resolve_template_path(identifier):
    """解析模板路径，支持文件路径和模板名称"""
    if os.path.isfile(identifier):
        return identifier
    
    current_dir = Path(__file__).parent
    candidates = [
        os.path.join("templates", identifier + ".json"),
        os.path.join("Chinese_paper_detect_templates", identifier + ".json"),
        os.path.join("services", "Chinese_paper_detect_templates", identifier + ".json"),
        str(current_dir.parent / "Chinese_paper_detect_templates" / (identifier + ".json")),
        str(current_dir / ".." / "Chinese_paper_detect_templates" / (identifier + ".json")),
    ]
    
    for candidate in candidates:
        abs_path = os.path.abspath(candidate)
        if os.path.isfile(abs_path):
            return abs_path
        if os.path.isfile(candidate):
            return candidate
    
    raise FileNotFoundError(f"Template not found: '{identifier}' (tried file path and {candidates})")

def load_template(identifier):
    """加载JSON模板文件"""
    tpl_path = resolve_template_path(identifier)
    with open(tpl_path, 'r', encoding='utf-8') as f:
        tpl = json.load(f)
    return tpl

# ---------- 字体检测函数 ----------
def detect_font_for_run(run, paragraph=None):
    """
    检测run的字体信息，包括中英文字体
    参数:
        run: docx run对象
        paragraph: docx paragraph对象
    返回: (font_size, font_ascii, font_eastasia, is_bold, is_italic, line_spacing)
    """
    font_size = None
    font_name_ascii = None
    font_name_eastasia = None
    is_bold = None
    is_italic = False
    
    if not run:
        return 12.0, "Times New Roman", "宋体", False, False, 1.0
    
    # 1. 检测字号
    try:
        if run.font and run.font.size and hasattr(run.font.size, 'pt'):
            font_size = float(run.font.size.pt)
        
        if font_size is None and getattr(run, "style", None) and getattr(run.style, "font", None):
            try:
                if run.style.font.size and hasattr(run.style.font.size, "pt"):
                    font_size = float(run.style.font.size.pt)
            except Exception:
                pass
        
        if font_size is None and paragraph and paragraph.style and getattr(paragraph.style, "font", None):
            try:
                if paragraph.style.font.size and hasattr(paragraph.style.font.size, "pt"):
                    font_size = float(paragraph.style.font.size.pt)
            except Exception:
                pass
        
        if font_size is None and paragraph and paragraph.style and hasattr(paragraph.style, "element"):
            sz_nodes = paragraph.style.element.xpath('.//w:sz')
            if sz_nodes and sz_nodes[0].get(qn('w:val')):
                font_size = float(sz_nodes[0].get(qn('w:val'))) / 2.0
        
        if font_size is None and hasattr(run._element, 'rPr'):
            sz_nodes = run._element.xpath('.//w:sz')
            if sz_nodes and sz_nodes[0].get(qn('w:val')):
                font_size = float(sz_nodes[0].get(qn('w:val'))) / 2.0
    except Exception:
        pass
    
    font_size = font_size if font_size is not None else 12.0
    
    # 2. 检测英文字体（ASCII）和中文字体（EastAsia）
    try:
        if run.font and run.font.name:
            font_name_ascii = run.font.name
        
        if hasattr(run._element, 'rPr'):
            rpr = run._element.rPr
            if rpr is not None:
                rfonts = rpr.find(qn('w:rFonts'))
                if rfonts is not None:
                    xml_ascii = rfonts.get(qn('w:ascii'))
                    xml_hansi = rfonts.get(qn('w:hAnsi'))
                    xml_eastasia = rfonts.get(qn('w:eastAsia'))
                    
                    if xml_ascii:
                        font_name_ascii = xml_ascii
                    elif xml_hansi and font_name_ascii is None:
                        font_name_ascii = xml_hansi
                    
                    if xml_eastasia:
                        font_name_eastasia = xml_eastasia
                    elif xml_hansi and font_name_eastasia is None:
                        font_name_eastasia = xml_hansi
    except Exception:
        pass
    
    font_name_ascii = font_name_ascii if font_name_ascii else "Times New Roman"
    font_name_eastasia = font_name_eastasia if font_name_eastasia else "宋体"
    
    # 3. 检测加粗和斜体
    try:
        if run.font:
            if run.font.bold is not None:
                is_bold = run.font.bold
            if run.font.italic is not None:
                is_italic = run.font.italic
        
        if is_bold is None and getattr(run, "style", None) and getattr(run.style, "font", None):
            try:
                if run.style.font.bold is not None:
                    is_bold = run.style.font.bold
            except Exception:
                pass
        
        if is_bold is None and paragraph and paragraph.style and getattr(paragraph.style, "font", None):
            try:
                if paragraph.style.font.bold is not None:
                    is_bold = paragraph.style.font.bold
            except Exception:
                pass
        
        if is_bold is None and hasattr(run._element, 'rPr'):
            try:
                b_nodes = run._element.xpath('.//w:b')
                if b_nodes:
                    is_bold = True
            except Exception:
                pass
        
        is_bold = bool(is_bold) if is_bold is not None else False
        is_italic = bool(is_italic)
    except Exception:
        pass
    
    # 4. 行间距检测
    line_spacing = 1.0
    try:
        if paragraph and paragraph.paragraph_format.line_spacing:
            line_spacing = float(paragraph.paragraph_format.line_spacing)
        elif paragraph and paragraph.style and paragraph.style.paragraph_format.line_spacing:
            line_spacing = float(paragraph.style.paragraph_format.line_spacing)
    except Exception:
        pass
    
    return font_size, font_name_ascii, font_name_eastasia, is_bold, is_italic, line_spacing

def get_font_size(pt_size, tpl=None):
    """字体大小转换为中文字号"""
    if tpl and 'check_rules' in tpl and 'font_size_mapping' in tpl['check_rules']:
        size_config = tpl['check_rules']['font_size_mapping']
        size_map = {}
        for key, value in size_config.items():
            try:
                size_map[float(key)] = value
            except (ValueError, TypeError):
                continue
    else:
        size_map = {
            9: "小五", 10.5: "五号", 12: "小四", 14: "四号",
            16: "三号", 18: "小二", 22: "二号", 24: "小一", 26: "一号"
        }
    
    if not size_map:
        return f"{pt_size}pt"
    
    closest_size = min(size_map.keys(), key=lambda x: abs(x - pt_size))
    return size_map[closest_size]

def get_line_spacing_name(spacing, tpl=None):
    """行间距转换为中文描述"""
    if tpl and 'check_rules' in tpl and 'line_spacing_mapping' in tpl['check_rules']:
        spacing_config = tpl['check_rules']['line_spacing_mapping']
        spacing_map = {}
        for key, value in spacing_config.items():
            try:
                spacing_map[float(key)] = value
            except (ValueError, TypeError):
                continue
    else:
        spacing_map = {
            1.0: "单倍行距", 1.15: "1.15倍行距", 
            1.25: "1.25倍行距", 1.5: "1.5倍行距", 2.0: "双倍行距"
        }
    
    if not spacing_map:
        return f"{spacing}倍行距"
    
    closest_spacing = min(spacing_map.keys(), key=lambda x: abs(x - spacing))
    return spacing_map[closest_spacing]

def get_alignment_name(alignment_value, tpl=None):
    """段落对齐方式转换为中文描述"""
    if tpl and 'check_rules' in tpl and 'alignment_mapping' in tpl['check_rules']:
        align_config = tpl['check_rules']['alignment_mapping']
        align_map = {}
        for key, value in align_config.items():
            align_map[key] = value
    else:
        align_map = {
            "left": "左对齐", "center": "居中对齐", 
            "right": "右对齐", "justify": "两端对齐"
        }
    
    alignment_map = {0: "左对齐", 1: "居中对齐", 2: "右对齐", 3: "两端对齐"}
    align_str = alignment_map.get(alignment_value, f"未知对齐({alignment_value})")
    
    # 如果模板中有映射，使用模板映射
    if align_map:
        # 尝试从模板中找到对应的描述
        for key, value in align_map.items():
            if value == align_str:
                return value
    
    return align_str

def detect_paragraph_alignment(paragraph):
    """检测段落对齐方式"""
    direct_alignment = paragraph.paragraph_format.alignment
    if direct_alignment is not None:
        return int(direct_alignment)
    
    if paragraph.style:
        try:
            style_alignment = paragraph.style.paragraph_format.alignment
            if style_alignment is not None:
                return int(style_alignment)
        except Exception:
            pass
    
    return WD_PARAGRAPH_ALIGNMENT.LEFT

def cm_to_pt(cm_value):
    """将厘米转换为磅（pt）"""
    return float(cm_value) * 28.35

# ---------- 结构检测 ----------
def find_toc_titles(doc, tpl):
    """
    查找目录、图录、表录标题
    返回: [{'title_type': '目录', 'paragraph': para, 'index': idx}, ...]
    """
    structure_rules = tpl.get('structure_rules', {})
    title_patterns = structure_rules.get('title_patterns', {})
    
    found_titles = []
    
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip() if paragraph.text else ''
        if not text:
            continue
        
        # 检查是否匹配某个标题模式
        for title_type, pattern in title_patterns.items():
            if re.match(pattern, text):
                found_titles.append({
                    'title_type': title_type,
                    'paragraph': paragraph,
                    'index': idx,
                    'text': text
                })
                break  # 每个段落只匹配一个标题
    
    return found_titles

def check_blank_lines_after_title(doc, title_info, tpl):
    """
    检查标题与条目之间的空行数
    参数:
        doc: Word文档对象
        title_info: 标题信息字典 {'title_type': str, 'paragraph': para, 'index': int}
        tpl: 模板配置
    返回: {'ok': bool, 'message': str}
    """
    structure_rules = tpl.get('structure_rules', {})
    expected_blank_lines = structure_rules.get('blank_lines_after_title', 2)
    
    title_idx = title_info['index']
    title_type = title_info['title_type']
    
    # 从标题后开始查找，跳过空段落
    blank_count = 0
    first_content_idx = None
    
    for i in range(title_idx + 1, len(doc.paragraphs)):
        para = doc.paragraphs[i]
        text = para.text.strip() if para.text else ''
        
        if not text:
            blank_count += 1
        else:
            first_content_idx = i
            break
    
    # 如果找到了内容，检查空行数
    if first_content_idx is not None:
        if blank_count == expected_blank_lines:
            msg_tpl = tpl.get('messages', {}).get('structure_blank_lines_ok', '{title_type}标题与条目之间空行正确（空两行）')
            return {'ok': True, 'message': msg_tpl.format(title_type=title_type)}
        else:
            msg_tpl = tpl.get('messages', {}).get('structure_blank_lines_error', '{title_type}标题与条目之间应空两行，实际为空{actual}行')
            return {'ok': False, 'message': msg_tpl.format(title_type=title_type, actual=blank_count)}
    else:
        # 没有找到内容，可能是文档末尾
        return {'ok': True, 'message': f'{title_type}标题后没有找到条目内容'}

def check_toc_structure(doc, tpl):
    """
    检查目录/图录/表录结构
    返回: {'ok': bool, 'messages': [], 'titles': [title_info, ...]}
    """
    report = {'ok': True, 'messages': [], 'titles': []}
    
    # 查找所有标题
    found_titles = find_toc_titles(doc, tpl)
    
    if not found_titles:
        report['ok'] = False
        # 检查每个可能的标题类型
        title_patterns = tpl.get('structure_rules', {}).get('title_patterns', {})
        for title_type in title_patterns.keys():
            msg_tpl = tpl.get('messages', {}).get('structure_title_not_found', '未找到{title_type}标题')
            report['messages'].append(msg_tpl.format(title_type=title_type))
        return report
    
    # 对每个找到的标题进行结构检查
    for title_info in found_titles:
        report['titles'].append(title_info)
        
        # 检查空行
        blank_check = check_blank_lines_after_title(doc, title_info, tpl)
        if not blank_check['ok']:
            report['ok'] = False
        report['messages'].append(blank_check['message'])
    
    return report

# ---------- 格式检测 ----------
def check_toc_format(paragraph, tpl, title_type):
    """
    检查目录/图录/表录标题格式
    返回: {'ok': bool, 'messages': []}
    """
    report = {'ok': True, 'messages': []}
    
    if not paragraph or not paragraph.runs:
        report['ok'] = False
        report['messages'].append(f"{title_type}标题段落没有文本内容")
        return report
    
    format_rules = tpl.get('format_rules', {}).get('title', {})
    issues = []
    
    # 检测第一个非空run的格式
    main_run = None
    for run in paragraph.runs:
        if run.text.strip():
            main_run = run
            break
    
    if not main_run:
        report['ok'] = False
        report['messages'].append(f"{title_type}标题段落没有有效文本")
        return report
    
    # 检测实际格式
    font_result = detect_font_for_run(main_run, paragraph)
    actual_size_pt, actual_font_ascii, actual_font_eastasia, actual_bold, actual_italic, actual_line_spacing = font_result
    
    # 字体大小检查
    if not should_skip_check('font_size') and 'font_size_pt' in format_rules:
        expected_size_pt = float(format_rules['font_size_pt'])
        actual_size_name = get_font_size(actual_size_pt, tpl)
        expected_size_name = get_font_size(expected_size_pt, tpl)
        if abs(actual_size_pt - expected_size_pt) > 0.5:
            issues.append(f"字体大小应为{expected_size_name}（{expected_size_pt}pt），实际为{actual_size_name}（{actual_size_pt}pt）")
    
    # 字体名称检查（中文字体）
    if not should_skip_check('font_name') and 'font_name' in format_rules:
        expected_font = format_rules.get('font_name', '黑体')
        if expected_font.lower() not in (actual_font_eastasia or '').lower():
            issues.append(f"字体应为{expected_font}，实际为{actual_font_eastasia}")
    
    # 加粗检查
    if not should_skip_check('bold') and 'bold' in format_rules:
        expected_bold = bool(format_rules['bold'])
        if actual_bold != expected_bold:
            bold_status = "加粗" if expected_bold else "不加粗"
            actual_status = "加粗" if actual_bold else "不加粗"
            issues.append(f"字体应为{bold_status}，实际为{actual_status}")
    
    # 行间距检查
    if not should_skip_check('spacing') and 'line_spacing' in format_rules:
        expected_line_spacing = float(format_rules['line_spacing'])
        if abs(actual_line_spacing - expected_line_spacing) > 0.1:
            actual_spacing_name = get_line_spacing_name(actual_line_spacing, tpl)
            expected_spacing_name = get_line_spacing_name(expected_line_spacing, tpl)
            issues.append(f"行间距应为{expected_spacing_name}（{expected_line_spacing}倍），实际为{actual_spacing_name}（{actual_line_spacing}倍）")
    
    # 段落对齐检查
    if 'alignment' in format_rules:
        expected_alignment_str = str(format_rules['alignment'])
        alignment_map = {"left": 0, "center": 1, "right": 2, "justify": 3}
        expected_alignment = alignment_map.get(expected_alignment_str, 0)
        
        actual_alignment = detect_paragraph_alignment(paragraph)
        if actual_alignment != expected_alignment:
            actual_alignment_name = get_alignment_name(actual_alignment, tpl)
            expected_alignment_name = get_alignment_name(expected_alignment, tpl)
            issues.append(f"段落应为{expected_alignment_name}，实际为{actual_alignment_name}")
    
    # 段前段后间距检查
    def pt(v): 
        return v.pt if v else 0.0
    
    if 'space_before_cm' in format_rules:
        space_before = paragraph.paragraph_format.space_before
        expected_before_pt = cm_to_pt(format_rules['space_before_cm'])
        actual_before_pt = pt(space_before)
        if abs(actual_before_pt - expected_before_pt) > 2.0:  # 2pt容差
            issues.append(f"段前间距应为{format_rules['space_before_cm']}厘米（约{expected_before_pt:.1f}pt），实际为{actual_before_pt:.1f}pt")
    
    if 'space_after_cm' in format_rules:
        space_after = paragraph.paragraph_format.space_after
        expected_after_pt = cm_to_pt(format_rules['space_after_cm'])
        actual_after_pt = pt(space_after)
        if abs(actual_after_pt - expected_after_pt) > 2.0:  # 2pt容差
            issues.append(f"段后间距应为{format_rules['space_after_cm']}厘米（约{expected_after_pt:.1f}pt），实际为{actual_after_pt:.1f}pt")
    
    if issues:
        report['ok'] = False
        header = tpl.get('messages', {}).get('format_toc_issue_header', '{title_type}格式问题：')
        report['messages'].append(header.format(title_type=title_type))
        report['messages'].extend([f"  - {i}" for i in issues])
    else:
        ok_msg = tpl.get('messages', {}).get('format_toc_ok', '{title_type}格式检查通过')
        report['messages'].append(ok_msg.format(title_type=title_type))
    
    return report

# ---------- 主入口 ----------
def check_toc_with_template(doc_path, template_identifier="TOC", skip_checks=None):
    """
    主检查函数：检查目录/图录/表录格式
    参数:
        doc_path: 文档路径
        template_identifier: 模板标识符，默认为 "TOC"
        skip_checks: 要跳过的检测项列表，如 ['font_size', 'bold']
    返回:
        检测结果字典
    """
    global _skip_checks_config
    _skip_checks_config = skip_checks or []
    
    tpl = load_template(template_identifier)
    doc = Document(doc_path)
    
    # 执行结构检查
    structure_report = check_toc_structure(doc, tpl)
    
    # 执行格式检查（对每个找到的标题）
    format_reports = []
    for title_info in structure_report.get('titles', []):
        format_report = check_toc_format(
            title_info['paragraph'], 
            tpl, 
            title_info['title_type']
        )
        format_reports.append({
            'title_type': title_info['title_type'],
            'report': format_report
        })
    
    # 组装报告
    overall_ok = structure_report.get('ok', True)
    for fr in format_reports:
        if not fr['report'].get('ok', True):
            overall_ok = False
    
    summary_tpl = tpl.get('messages', {}).get('summary_overall', "目录/图录/表录检查结果: {ok}")
    summary = summary_tpl.format(ok="通过" if overall_ok else "失败")
    
    report = {
        "ok": overall_ok,
        "structure": structure_report,
        "format": format_reports,
        "summary": [summary]
    }
    
    return report

# ---------- 命令行入口 ----------
def main():
    import argparse
    parser = argparse.ArgumentParser(description='检测目录/图录/表录格式')
    parser.add_argument('doc_path', help='Word文档路径')
    parser.add_argument('-t', '--template', default='TOC', help='模板标识符（默认: TOC）')
    parser.add_argument('--skip', nargs='+', help='跳过的检测项，如: --skip font_size bold')
    
    args = parser.parse_args()
    
    skip_checks = args.skip if args.skip else None
    result = check_toc_with_template(args.doc_path, args.template, skip_checks)
    
    print("=" * 80)
    print("【目录/图录/表录 检测报告】")
    print("=" * 80)
    print()
    
    # 结构检测报告
    print("  [Structure]", "✓ 通过" if result['structure']['ok'] else "✗ 失败")
    for msg in result['structure']['messages']:
        print(f"    • {msg}")
    print()
    
    # 格式检测报告
    for fr in result['format']:
        title_type = fr['title_type']
        format_report = fr['report']
        print(f"  [{title_type} Format]", "✓ 通过" if format_report['ok'] else "✗ 失败")
        for msg in format_report['messages']:
            print(f"    • {msg}")
        print()
    
    # 总结
    print("  【总结】")
    for s in result['summary']:
        print(f"    {s}")
    print()

if __name__ == "__main__":
    main()

