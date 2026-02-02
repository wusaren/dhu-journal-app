#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""调试脚本：检查文档中的二级标题"""
import sys
import re
from docx import Document

def debug_level2_titles(doc_path):
    doc = Document(doc_path)
    level2_pattern = r'^\s*(\d+\.\d+)\s+(.+)$'
    
    print("=" * 80)
    print("检查文档中的二级标题")
    print("=" * 80)
    
    potential_level2 = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip() if paragraph.text else ''
        if not text:
            continue
        
        # 检查是否匹配二级标题模式
        if re.match(level2_pattern, text):
            style_name = getattr(paragraph.style, 'name', '') if paragraph.style else ''
            potential_level2.append({
                'index': idx,
                'text': text,
                'style': style_name,
                'full_text': paragraph.text
            })
    
    print(f"\n找到 {len(potential_level2)} 个匹配二级标题模式的段落：\n")
    
    for item in potential_level2:
        print(f"段落索引: {item['index']}")
        print(f"样式名称: {item['style']}")
        print(f"段落文本: {item['text']}")
        print(f"完整文本: {repr(item['full_text'])}")
        # 检查是否是toc样式
        if 'toc' in item['style'].lower():
            print("⚠️  这是目录中的标题（toc样式），会被过滤掉")
        else:
            print("✓ 这是正文中的标题，会被检测")
        print("-" * 80)
    
    # 检查所有包含数字.数字格式的段落
    print("\n" + "=" * 80)
    print("检查所有包含 '数字.数字' 格式的段落（可能是二级标题）")
    print("=" * 80)
    
    all_potential = []
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip() if paragraph.text else ''
        if not text:
            continue
        
        # 检查是否包含数字.数字格式
        if re.search(r'\d+\.\d+', text):
            style_name = getattr(paragraph.style, 'name', '') if paragraph.style else ''
            match = re.match(level2_pattern, text)
            is_match = match is not None
            all_potential.append({
                'index': idx,
                'text': text[:60] + ('...' if len(text) > 60 else ''),
                'style': style_name,
                'is_match': is_match,
                'full_text': paragraph.text[:80] + ('...' if len(paragraph.text) > 80 else '')
            })
    
    print(f"\n找到 {len(all_potential)} 个包含 '数字.数字' 格式的段落：\n")
    
    for item in all_potential[:20]:  # 只显示前20个
        match_status = "[匹配]" if item['is_match'] else "[不匹配]"
        print(f"[{item['index']}] {match_status} | 样式: {item['style']:15} | {item['text']}")
    
    if len(all_potential) > 20:
        print(f"\n... 还有 {len(all_potential) - 20} 个段落未显示")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("用法: python debug_level2.py <docx文件路径>")
        sys.exit(1)
    
    doc_path = sys.argv[1]
    debug_level2_titles(doc_path)

