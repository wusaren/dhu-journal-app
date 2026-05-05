"""
批量检测Excel导出服务
用于导出业务员批量检测的论文审核情况汇总表
"""

import os
import json
import logging
from datetime import datetime
from typing import Dict, Any, List, Optional

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from models import BatchJob, PaperCheckResult

logger = logging.getLogger(__name__)


class BatchExcelExportService:
    """批量检测Excel导出服务"""
    
    def __init__(self):
        # 样式定义
        self.header_font = Font(name='宋体', size=12, bold=True, color='FFFFFF')
        self.header_fill = PatternFill(start_color='9C0E0E', end_color='9C0E0E', fill_type='solid')
        self.header_alignment = Alignment(horizontal='center', vertical='center')
        
        self.cell_font = Font(name='宋体', size=11)
        self.cell_alignment = Alignment(horizontal='center', vertical='center')
        
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 状态颜色
        self.pass_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')  # 绿色
        self.fail_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')  # 红色
        self.pending_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')  # 黄色
        self.error_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')  # 灰色
    
    def export_batch_excel(self, batch_id: int) -> Dict[str, Any]:
        """
        导出批次Excel汇总表
        
        Args:
            batch_id: 批次ID
            
        Returns:
            {'success': bool, 'file_path': str, 'filename': str, 'message': str}
        """
        try:
            batch_job = BatchJob.query.get(batch_id)
            if not batch_job:
                return {'success': False, 'message': '批次任务不存在'}
            
            papers = PaperCheckResult.query.filter_by(batch_job_id=batch_id).order_by(
                PaperCheckResult.student_id.asc()
            ).all()
            
            if not papers:
                return {'success': False, 'message': '批次中没有论文数据'}
            
            # 创建Excel工作簿
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = '论文审核汇总'
            
            # 设置列宽
            column_widths = {
                'A': 8,   # 序号
                'B': 12,  # 学号
                'C': 10,  # 姓名
                'D': 40,  # 文件名
                'E': 12,  # 检测状态
                'F': 12,  # 审核状态
                'G': 10,  # 通过率
                'H': 20,  # 检测时间
                'I': 50,  # 批注文件
                'J': 50,  # 报告文件
                'K': 100, # 错误信息
            }
            for col, width in column_widths.items():
                ws.column_dimensions[col].width = width
            
            # 写入标题行
            headers = ['序号', '学号', '姓名', '文件名', '检测状态', '审核状态', '通过率', '检测时间', '批注文件', '报告文件', '错误信息']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = self.header_font
                cell.fill = self.header_fill
                cell.alignment = self.header_alignment
                cell.border = self.border
            
            # 设置标题行高
            ws.row_dimensions[1].height = 30
            
            # 写入数据行
            for row, paper in enumerate(papers, 2):
                # 序号
                ws.cell(row=row, column=1, value=row-1).font = self.cell_font
                ws.cell(row=row, column=1).alignment = self.cell_alignment
                ws.cell(row=row, column=1).border = self.border
                
                # 学号
                ws.cell(row=row, column=2, value=paper.student_id or '').font = self.cell_font
                ws.cell(row=row, column=2).alignment = self.cell_alignment
                ws.cell(row=row, column=2).border = self.border
                
                # 姓名
                ws.cell(row=row, column=3, value=paper.student_name or '').font = self.cell_font
                ws.cell(row=row, column=3).alignment = self.cell_alignment
                ws.cell(row=row, column=3).border = self.border
                
                # 文件名
                ws.cell(row=row, column=4, value=paper.original_filename or '').font = self.cell_font
                ws.cell(row=row, column=4).alignment = Alignment(horizontal='left', vertical='center')
                ws.cell(row=row, column=4).border = self.border
                
                # 检测状态
                check_status_map = {
                    'pending': '待检测',
                    'completed': '已完成',
                    'failed': '失败'
                }
                check_status_text = check_status_map.get(paper.check_status, paper.check_status or '')
                cell_check = ws.cell(row=row, column=5, value=check_status_text)
                cell_check.font = self.cell_font
                cell_check.alignment = self.cell_alignment
                cell_check.border = self.border
                # 设置检测状态颜色
                if paper.check_status == 'completed':
                    cell_check.fill = self.pass_fill
                elif paper.check_status == 'failed':
                    cell_check.fill = self.error_fill
                else:
                    cell_check.fill = self.pending_fill
                
                # 审核状态
                review_status_map = {
                    'pending': '待审核',
                    'reviewed': '已通过',
                    'needs_revision': '需修改'
                }
                review_status_text = review_status_map.get(paper.review_status, paper.review_status or '')
                cell_review = ws.cell(row=row, column=6, value=review_status_text)
                cell_review.font = self.cell_font
                cell_review.alignment = self.cell_alignment
                cell_review.border = self.border
                # 设置审核状态颜色
                if paper.review_status == 'reviewed':
                    cell_review.fill = self.pass_fill
                elif paper.review_status == 'needs_revision':
                    cell_review.fill = self.fail_fill
                else:
                    cell_review.fill = self.pending_fill
                
                # 通过率
                pass_rate_text = f'{paper.pass_rate:.1f}%' if paper.pass_rate is not None else 'N/A'
                cell_rate = ws.cell(row=row, column=7, value=pass_rate_text)
                cell_rate.font = self.cell_font
                cell_rate.alignment = self.cell_alignment
                cell_rate.border = self.border
                
                # 检测时间
                completed_time = paper.completed_at.strftime('%Y-%m-%d %H:%M') if paper.completed_at else 'N/A'
                ws.cell(row=row, column=8, value=completed_time).font = self.cell_font
                ws.cell(row=row, column=8).alignment = self.cell_alignment
                ws.cell(row=row, column=8).border = self.border
                
                # 批注文件路径
                annotated_name = os.path.basename(paper.annotated_path) if paper.annotated_path else 'N/A'
                ws.cell(row=row, column=9, value=annotated_name).font = self.cell_font
                ws.cell(row=row, column=9).alignment = Alignment(horizontal='left', vertical='center')
                ws.cell(row=row, column=9).border = self.border
                
                # 报告文件路径
                report_name = os.path.basename(paper.report_path) if paper.report_path else 'N/A'
                ws.cell(row=row, column=10, value=report_name).font = self.cell_font
                ws.cell(row=row, column=10).alignment = Alignment(horizontal='left', vertical='center')
                ws.cell(row=row, column=10).border = self.border
                
                # 错误信息
                ws.cell(row=row, column=11, value=paper.error_message or '').font = self.cell_font
                ws.cell(row=row, column=11).alignment = Alignment(horizontal='left', vertical='center')
                ws.cell(row=row, column=11).border = self.border
                
                # 设置数据行高
                ws.row_dimensions[row].height = 25
            
            # 添加统计行
            stats_row = len(papers) + 3
            ws.cell(row=stats_row, column=1, value='统计汇总').font = Font(name='宋体', size=12, bold=True)
            ws.cell(row=stats_row+1, column=1, value=f'总论文数: {batch_job.total_papers}').font = self.cell_font
            ws.cell(row=stats_row+2, column=1, value=f'已检测: {batch_job.processed}').font = self.cell_font
            ws.cell(row=stats_row+3, column=1, value=f'通过: {batch_job.passed}').font = Font(name='宋体', size=11, color='006600')
            ws.cell(row=stats_row+4, column=1, value=f'需修改: {sum(1 for p in papers if p.review_status == "needs_revision")}').font = Font(name='宋体', size=11, color='CC0000')
            ws.cell(row=stats_row+5, column=1, value=f'失败: {batch_job.failed}').font = Font(name='宋体', size=11, color='666666')
            
            # 添加批次信息
            info_row = stats_row + 7
            ws.cell(row=info_row, column=1, value='批次信息').font = Font(name='宋体', size=12, bold=True)
            ws.cell(row=info_row+1, column=1, value=f'批次ID: {batch_job.id}').font = self.cell_font
            ws.cell(row=info_row+2, column=1, value=f'源目录: {batch_job.directory_path}').font = self.cell_font
            ws.cell(row=info_row+3, column=1, value=f'输出目录: {batch_job.output_path}').font = self.cell_font
            ws.cell(row=info_row+4, column=1, value=f'自动通过阈值: {batch_job.pass_threshold}%').font = self.cell_font
            ws.cell(row=info_row+5, column=1, value=f'开始时间: {batch_job.started_at.strftime("%Y-%m-%d %H:%M:%S") if batch_job.started_at else "N/A"}').font = self.cell_font
            ws.cell(row=info_row+6, column=1, value=f'完成时间: {batch_job.completed_at.strftime("%Y-%m-%d %H:%M:%S") if batch_job.completed_at else "N/A"}').font = self.cell_font
            
            # 保存文件
            date_str = datetime.now().strftime('%Y%m%d')
            filename = f'汇总统计表_{date_str}_batch{batch_id}.xlsx'
            file_path = os.path.join(batch_job.output_path, filename)
            
            wb.save(file_path)
            
            logger.info(f"Excel汇总表已生成: {file_path}")
            
            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'message': f'Excel汇总表已生成，共 {len(papers)} 条记录'
            }
            
        except Exception as e:
            logger.error(f"导出Excel失败: {e}")
            return {'success': False, 'message': f'导出Excel失败: {str(e)}'}
    
    def get_batch_summary(self, batch_id: int) -> Optional[Dict[str, Any]]:
        """获取批次汇总信息"""
        try:
            batch_job = BatchJob.query.get(batch_id)
            if not batch_job:
                return None
            
            papers = PaperCheckResult.query.filter_by(batch_job_id=batch_id).all()
            
            # 统计审核状态
            reviewed_count = sum(1 for p in papers if p.review_status == 'reviewed')
            needs_revision_count = sum(1 for p in papers if p.review_status == 'needs_revision')
            pending_count = sum(1 for p in papers if p.review_status == 'pending')
            
            # 计算平均通过率
            completed_papers = [p for p in papers if p.pass_rate is not None]
            avg_pass_rate = sum(p.pass_rate for p in completed_papers) / len(completed_papers) if completed_papers else 0
            
            return {
                'batch_id': batch_id,
                'total_papers': batch_job.total_papers,
                'processed': batch_job.processed,
                'passed': batch_job.passed,
                'failed': batch_job.failed,
                'reviewed_count': reviewed_count,
                'needs_revision_count': needs_revision_count,
                'pending_count': pending_count,
                'avg_pass_rate': round(avg_pass_rate, 2),
                'status': batch_job.status,
                'pass_threshold': batch_job.pass_threshold
            }
            
        except Exception as e:
            logger.error(f"获取批次汇总失败: {e}")
            return None
