#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
=== 论文格式检测系统 - 参考文献引用检测器 ===

【参考文献引用检测 (Reference Citation Detection)】

1. 【参考文献列表识别】
   - 识别"References"或"参考文献"标题
   - 提取所有参考文献条目（[1], [2], ...）

2. 【正文引用识别】
   - 支持单个引用：[1]
   - 支持范围引用：[1-3] → [1, 2, 3]
   - 支持多个引用：[1,3,5] → [1, 3, 5]
   - 支持混合引用：[1-3,5-7] → [1, 2, 3, 5, 6, 7]

3. 【引用顺序检查】
   - 文中引用应按递增顺序出现
   - 报告乱序的位置

4. 【未引用检查】
   - 检查哪些参考文献未被引用

5. 【不存在引用检查】
   - 检查引用了但不在参考文献列表中的编号
"""

import os
import sys
import json
import re
from docx import Document

# 全局检测配置（由 run_all_detections 在导入时注入）
GLOBAL_DETECTION_CONFIG = {'skip_checks': set()}

def should_skip_check(check_name):
    """判断是否应该跳过某个检测项"""
    return check_name in GLOBAL_DETECTION_CONFIG.get('skip_checks', set())

# ---------- 模板加载 ----------
def resolve_template_path(identifier):
    """解析模板路径，支持文件路径和模板名称"""
    if os.path.isfile(identifier):
        return identifier
    candidate = os.path.join("templates", identifier + ".json")
    if os.path.isfile(candidate):
        return candidate
    raise FileNotFoundError(f"Template not found: '{identifier}' (tried file path and {candidate})")

def load_template(identifier):
    """加载JSON模板文件"""
    tpl_path = resolve_template_path(identifier)
    with open(tpl_path, 'r', encoding='utf-8') as f:
        tpl = json.load(f)
    return tpl

# ---------- 参考文献识别 ----------
def find_references_section(doc):
    """
    查找参考文献部分的位置
    
    返回：
        (start_index, end_index) 或 None
        start_index: 参考文献标题的段落索引
        end_index: 参考文献部分结束的段落索引
    """
    ref_start = None
    
    # 查找参考文献标题
    for idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if re.match(r'^(References|参考文献|REFERENCES)\s*$', text, re.IGNORECASE):
            ref_start = idx
            break
    
    if ref_start is None:
        return None
    
    # 查找参考文献结束位置（遇到中文标题或文档结束）
    ref_end = len(doc.paragraphs)
    for idx in range(ref_start + 1, len(doc.paragraphs)):
        text = doc.paragraphs[idx].text.strip()
        # 检测到中文标题（中文部分开始）
        if re.match(r'^[\u4e00-\u9fa5]', text) and len(text) > 5:
            ref_end = idx
            break
    
    return (ref_start, ref_end)

def extract_references(doc, start_index, end_index):
    """
    提取参考文献列表
    
    返回：
        [{'number': int, 'text': str, 'paragraph_index': int}, ...]
    """
    references = []
    
    for idx in range(start_index + 1, end_index):
        text = doc.paragraphs[idx].text.strip()
        
        if not text:
            continue
        
        # 匹配参考文献格式：[1], [2], ...
        match = re.match(r'^\[(\d+)\]\s*(.+)', text)
        if match:
            ref_num = int(match.group(1))
            ref_text = match.group(2)
            references.append({
                'number': ref_num,
                'text': ref_text,
                'paragraph_index': idx
            })
    
    return references

# ---------- 正文引用识别 ----------
def parse_citation(citation_str):
    """
    解析引用字符串，返回所有引用的编号列表
    
    例如：
    - "1" -> [1]
    - "1-3" -> [1, 2, 3]
    - "1,3,5" -> [1, 3, 5]
    - "1-3,5-7" -> [1, 2, 3, 5, 6, 7]
    """
    numbers = []
    parts = citation_str.split(',')
    
    for part in parts:
        part = part.strip()
        if '-' in part:
            # 范围引用
            try:
                start, end = part.split('-')
                start = int(start.strip())
                end = int(end.strip())
                numbers.extend(range(start, end + 1))
            except ValueError:
                continue
        else:
            # 单个引用
            try:
                numbers.append(int(part))
            except ValueError:
                continue
    
    return numbers

def find_citations_in_text(doc, ref_start_index):
    """
    查找正文中的所有引用
    
    返回：
        [{'paragraph_index': int, 'citation_str': str, 'numbers': [int], 'position': int}, ...]
    """
    citations = []
    
    # 只分析参考文献之前的内容
    for idx in range(0, ref_start_index):
        text = doc.paragraphs[idx].text
        
        # 查找引用模式：[1], [2-5], [1,3,5], [1-3,5-7]
        citation_pattern = r'\[(\d+(?:-\d+)?(?:,\s*\d+(?:-\d+)?)*)\]'
        matches = re.finditer(citation_pattern, text)
        
        for match in matches:
            citation_str = match.group(1)
            cited_nums = parse_citation(citation_str)
            
            if cited_nums:
                citations.append({
                    'paragraph_index': idx,
                    'citation_str': citation_str,
                    'numbers': cited_nums,
                    'position': match.start()
                })
    
    return citations

# ---------- 引用检查 ----------
def check_citation_order(citations):
    """
    检查引用顺序是否正确（递增）
    
    返回：
        {'ok': bool, 'errors': []}
    """
    report = {'ok': True, 'errors': []}
    
    # 提取所有引用的编号（按出现顺序）
    all_cited_nums = []
    citation_positions = []  # 记录每个编号的位置信息
    
    for citation in citations:
        for num in citation['numbers']:
            all_cited_nums.append(num)
            citation_positions.append({
                'number': num,
                'paragraph_index': citation['paragraph_index'],
                'citation_str': citation['citation_str']
            })
    
    # 检查是否按顺序引用
    for i in range(len(all_cited_nums) - 1):
        if all_cited_nums[i] > all_cited_nums[i + 1]:
            report['ok'] = False
            pos1 = citation_positions[i]
            pos2 = citation_positions[i + 1]
            
            report['errors'].append({
                'type': 'warning',
                'page_number': f"段落{pos2['paragraph_index']}",
                'description': f"引用顺序不正确：[{pos1['number']}]后出现了[{pos2['number']}]（应该递增）",
                'suggestion': f"建议：调整引用顺序，确保引用编号按递增顺序出现",
                'text_snippet': f"[{pos2['citation_str']}]"  # 只使用第二个引用的文本，更容易定位
            })
    
    return report

def check_uncited_references(references, citations):
    """
    检查未被引用的参考文献
    
    返回：
        {'ok': bool, 'errors': []}
    """
    report = {'ok': True, 'errors': []}
    
    # 提取所有被引用的编号
    cited_nums = set()
    for citation in citations:
        cited_nums.update(citation['numbers'])
    
    # 检查哪些参考文献未被引用
    ref_nums = [ref['number'] for ref in references]
    uncited = [num for num in ref_nums if num not in cited_nums]
    
    if uncited:
        report['ok'] = False
        uncited_str = ', '.join(f'[{num}]' for num in sorted(uncited))
        
        # 为每个未引用的参考文献生成一个错误
        for num in sorted(uncited):
            ref_info = next((r for r in references if r['number'] == num), None)
            if ref_info:
                report['errors'].append({
                    'type': 'warning',
                    'page_number': f"段落{ref_info['paragraph_index']}",
                    'description': f"参考文献[{num}]未在正文中被引用",
                    'suggestion': '建议：在正文中添加对该参考文献的引用，或删除未使用的参考文献',
                    'text_snippet': ref_info['text'][:100] + ('...' if len(ref_info['text']) > 100 else '')
                })
    
    return report

def check_nonexistent_citations(references, citations):
    """
    检查引用了但不存在的参考文献
    
    返回：
        {'ok': bool, 'errors': []}
    """
    report = {'ok': True, 'errors': []}
    
    # 提取所有参考文献编号
    ref_nums = set(ref['number'] for ref in references)
    
    # 检查哪些引用不存在
    cited_nums = set()
    for citation in citations:
        cited_nums.update(citation['numbers'])
    
    nonexistent = sorted(cited_nums - ref_nums)
    
    if nonexistent:
        report['ok'] = False
        
        # 找出每个不存在引用的位置
        for num in nonexistent:
            # 找到第一次引用该编号的位置
            for citation in citations:
                if num in citation['numbers']:
                    report['errors'].append({
                        'type': 'error',
                        'page_number': f"段落{citation['paragraph_index']}",
                        'description': f"引用了不存在的参考文献[{num}]",
                        'suggestion': f"建议：添加参考文献[{num}]到参考文献列表，或删除该引用",
                        'text_snippet': f"[{citation['citation_str']}]"
                    })
                    break
    
    return report

# ---------- 主检查函数 ----------
def check_doc_with_template(doc_path, template_identifier):
    """
    主检查函数：使用模板检查文档中的参考文献引用
    
    返回：
        {'ok': bool, 'errors': []}
    """
    tpl = load_template(template_identifier)
    doc = Document(doc_path)
    
    report = {'ok': True, 'errors': []}
    
    # 1. 查找参考文献部分
    ref_section = find_references_section(doc)
    
    if ref_section is None:
        report['ok'] = False
        report['errors'].append({
            'type': 'error',
            'page_number': 'N/A',
            'description': '未找到参考文献部分',
            'suggestion': '建议：添加"References"或"参考文献"标题及参考文献列表',
            'text_snippet': 'N/A'
        })
        return report
    
    ref_start, ref_end = ref_section
    
    # 2. 提取参考文献列表
    references = extract_references(doc, ref_start, ref_end)
    
    if not references:
        report['ok'] = False
        report['errors'].append({
            'type': 'error',
            'page_number': f"段落{ref_start}",
            'description': '参考文献列表为空',
            'suggestion': '建议：添加参考文献条目，格式为[1] 作者. 标题[J]. 期刊, 年份, ...',
            'text_snippet': 'N/A'
        })
        return report
    
    # 3. 查找正文中的引用
    citations = find_citations_in_text(doc, ref_start)
    
    if not citations:
        report['ok'] = False
        report['errors'].append({
            'type': 'warning',
            'page_number': 'N/A',
            'description': '正文中未找到任何参考文献引用',
            'suggestion': '建议：在正文中添加对参考文献的引用，格式为[1]、[1-3]等',
            'text_snippet': 'N/A'
        })
        return report
    
    # 4. 检查引用顺序
    if not should_skip_check('citation_order'):
        order_report = check_citation_order(citations)
        if not order_report['ok']:
            report['ok'] = False
            report['errors'].extend(order_report['errors'])
    
    # 5. 检查未引用的参考文献
    if not should_skip_check('uncited_references'):
        uncited_report = check_uncited_references(references, citations)
        if not uncited_report['ok']:
            report['ok'] = False
            report['errors'].extend(uncited_report['errors'])
    
    # 6. 检查引用了不存在的参考文献
    if not should_skip_check('nonexistent_citations'):
        nonexistent_report = check_nonexistent_citations(references, citations)
        if not nonexistent_report['ok']:
            report['ok'] = False
            report['errors'].extend(nonexistent_report['errors'])
    
    return report

# ---------- 命令行接口 ----------
def print_help():
    """显示帮助信息"""
    print("用法:")
    print("  python Reference_detect.py check <paper.docx> <template.json_or_name>")
    print("")
    print("示例:")
    print("  python Reference_detect.py check template/test.docx Reference")
    print("  python Reference_detect.py check template/test.docx templates/Reference.json")

def main():
    """命令行入口"""
    if len(sys.argv) < 4 or sys.argv[1] != 'check':
        print_help()
        sys.exit(1)
    
    doc_path = sys.argv[2]
    template_identifier = sys.argv[3]
    
    if not os.path.isfile(doc_path):
        print(f"错误：文件不存在: {doc_path}")
        sys.exit(1)
    
    print("=" * 60)
    print("参考文献引用检测")
    print("=" * 60)
    print(f"文档: {doc_path}")
    print(f"模板: {template_identifier}")
    print()
    
    try:
        report = check_doc_with_template(doc_path, template_identifier)
        
        print(f"检测结果: {'✓ 通过' if report['ok'] else '✗ 发现问题'}")
        print(f"错误数量: {len(report.get('errors', []))}")
        print()
        
        if report.get('errors'):
            print("错误详情:")
            print("-" * 60)
            for i, error in enumerate(report['errors'], 1):
                print(f"\n错误 {i}:")
                print(f"  类型: {error.get('type')}")
                print(f"  位置: {error.get('page_number')}")
                print(f"  描述: {error.get('description')}")
                print(f"  建议: {error.get('suggestion')}")
                snippet = error.get('text_snippet', '')
                if snippet and snippet != 'N/A':
                    print(f"  文本: {snippet[:80]}...")
        
        print()
        print("=" * 60)
        
    except Exception as e:
        print(f"✗ 检测失败: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

if __name__ == '__main__':
    main()
